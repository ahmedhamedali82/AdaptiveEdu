#!/usr/bin/env python3
"""AdaptiveEdu project plan — Telegram study bot + two Google Sheets."""

from pathlib import Path
import shutil

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path("/Users/ahmedhamed/Documents/AI Agent/Technical/Graduation_Project/AdaptiveEdu")
ASSETS = Path("/Users/ahmedhamed/.cursor/projects/Users-ahmedhamed-Documents-AI-Agent/assets")
OUT = ROOT / "AdaptiveEdu_Project_Plan.pptx"
DIAG = ROOT / "diagrams"
DIAG.mkdir(parents=True, exist_ok=True)

W, H = Inches(13.333), Inches(7.5)
NAVY = RGBColor(0x0F, 0x27, 0x44)
NAVY2 = RGBColor(0x1B, 0x3A, 0x61)
TEAL = RGBColor(0x0E, 0x8C, 0x7E)
GOLD = RGBColor(0xE6, 0x9B, 0x2D)
ORANGE = RGBColor(0xE8, 0x78, 0x30)
GREEN = RGBColor(0x2E, 0x8C, 0x5A)
RED = RGBColor(0xB9, 0x1C, 0x1C)
LIGHT = RGBColor(0xF6, 0xF8, 0xFB)
SOFT_TEAL = RGBColor(0xE4, 0xF4, 0xF1)
SOFT_BLUE = RGBColor(0xE6, 0xEE, 0xF8)
SOFT_GOLD = RGBColor(0xFF, 0xF6, 0xE4)
SOFT_GREEN = RGBColor(0xE6, 0xF5, 0xEC)
SOFT_ORANGE = RGBColor(0xFF, 0xF0, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1C, 0x24, 0x30)
MUTED = RGBColor(0x5A, 0x66, 0x76)
PAGE = {"n": 0}

SHEET_STUDENTS = "https://docs.google.com/spreadsheets/d/17A-hQ6jQ7QJrcZ4f5f02O82PKKAtQ2gTiJHtv71LWI4"
SHEET_CURR = "https://docs.google.com/spreadsheets/d/10JNorDKgaesXGobuq_9sHipzRp-DOjxGAFZEV5D7Hwk"
GITHUB = "https://github.com/ahmedhamedali82/AdaptiveEdu"


def rgb_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_rect(slide, l, t, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    rgb_fill(sh, color)
    return sh


def add_round(slide, l, t, w, h, color, adj=0.08):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    rgb_fill(sh, color)
    try:
        sh.adjustments[0] = adj
    except Exception:
        pass
    return sh


def add_text(slide, l, t, w, h, text, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return box


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def chrome(slide, footer):
    add_rect(slide, 0, 0, W, H, WHITE)
    add_rect(slide, 0, 0, Inches(0.12), H, NAVY)
    add_rect(slide, Inches(0.12), Inches(7.22), Inches(13.213), Inches(0.28), LIGHT)
    add_text(slide, Inches(0.35), Inches(7.22), Inches(10.5), Inches(0.28), footer, size=11, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    PAGE["n"] += 1
    add_text(slide, Inches(11.5), Inches(7.22), Inches(1.5), Inches(0.28), str(PAGE["n"]), size=11, color=MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def header(slide, kicker, title, footer):
    chrome(slide, footer)
    add_text(slide, Inches(0.4), Inches(0.16), Inches(12.4), Inches(0.26), kicker, size=13, bold=True, color=TEAL)
    add_text(slide, Inches(0.4), Inches(0.40), Inches(12.4), Inches(0.44), title, size=24, bold=True, color=NAVY2)
    add_rect(slide, Inches(0.4), Inches(0.88), Inches(1.4), Inches(0.06), NAVY)


def card(slide, l, t, w, h, fill, title, body, bar=None, tsize=16, bsize=14):
    add_round(slide, l, t, w, h, fill)
    if bar:
        add_rect(slide, l, t, Inches(0.1), h, bar)
    add_text(slide, l + Inches(0.2), t + Inches(0.14), w - Inches(0.35), Inches(0.4), title, size=tsize, bold=True, color=NAVY)
    add_text(slide, l + Inches(0.2), t + Inches(0.55), w - Inches(0.35), h - Inches(0.7), body, size=bsize, color=DARK)


def pic(slide, name, left, top, width):
    for p in (DIAG / name, ASSETS / name, ASSETS / name.replace(".png", ".jpg")):
        if p.exists():
            slide.shapes.add_picture(str(p), left, top, width=width)
            return


def copy_imgs():
    pairs = [
        ("adaptiveedu_two_sheets.jpg", "adaptiveedu_two_sheets.jpg"),
        ("adaptiveedu_milestones_2weeks.jpg", "adaptiveedu_milestones.jpg"),
        ("adaptiveedu_milestones.jpg", "adaptiveedu_milestones.jpg"),
    ]
    for src, dest in pairs:
        sp = ASSETS / src
        if sp.exists() and dest == "adaptiveedu_milestones.jpg" and src.endswith("2weeks.jpg"):
            shutil.copy2(sp, DIAG / dest)
            break
    for src, dest in [
        ("adaptiveedu_two_sheets.jpg", "adaptiveedu_two_sheets.jpg"),
    ]:
        sp = ASSETS / src
        if sp.exists():
            shutil.copy2(sp, DIAG / dest)
    # 2-week timeline if present
    two = ASSETS / "adaptiveedu_milestones_2weeks.jpg"
    if two.exists():
        shutil.copy2(two, DIAG / "adaptiveedu_milestones.jpg")
    elif (ASSETS / "adaptiveedu_milestones.jpg").exists() and not (DIAG / "adaptiveedu_milestones.jpg").exists():
        shutil.copy2(ASSETS / "adaptiveedu_milestones.jpg", DIAG / "adaptiveedu_milestones.jpg")


def build():
    PAGE["n"] = 0
    copy_imgs()
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    f = "AdaptiveEdu  ·  Graduation project plan"

    # 1 Cover
    s = blank(prs)
    add_rect(s, 0, 0, W, H, NAVY)
    add_rect(s, Inches(8.6), 0, Inches(4.733), H, NAVY2)
    add_text(s, Inches(0.55), Inches(1.35), Inches(7.6), Inches(0.32), "DEPI GRADUATION PROJECT", size=14, bold=True, color=TEAL)
    add_text(s, Inches(0.55), Inches(1.85), Inches(7.6), Inches(1.0), "AdaptiveEdu", size=44, bold=True, color=WHITE)
    add_text(
        s, Inches(0.55), Inches(3.15), Inches(7.6), Inches(1.6),
        "A Telegram study bot for 1st Secondary.\nThe next lesson. A short quiz.\nPass to continue. Fail to retry simpler.",
        size=20, color=WHITE,
    )
    add_text(s, Inches(0.55), Inches(5.1), Inches(7.6), Inches(0.7), "n8n  ·  two Google Sheets  ·  Telegram", size=16, color=WHITE)
    toc = ["Problem", "How it works", "Two Google Sheets", "n8n workflow", "Pass / Fail rule", "2-week plan", "Deliverables"]
    for i, line in enumerate(toc):
        add_text(s, Inches(8.9), Inches(1.55 + i * 0.65), Inches(4.1), Inches(0.5), f"{i+1}  {line}", size=16, color=WHITE)
    PAGE["n"] += 1
    add_text(s, Inches(12.2), Inches(7.15), Inches(0.9), Inches(0.28), "1", size=11, color=WHITE, align=PP_ALIGN.RIGHT)

    # 2 Problem + solution
    s = blank(prs)
    header(s, "WHY THIS PROJECT", "Problem and solution", f)
    card(
        s, Inches(0.4), Inches(1.15), Inches(6.15), Inches(5.7), SOFT_ORANGE,
        "Problem",
        "One class. One pace. One book.\n\nA student who fails a lesson still moves on.\nA teacher cannot follow every student by hand.\n\nWe need a simple path:\nfinish this lesson before the next one.",
        bar=ORANGE, tsize=20, bsize=17,
    )
    card(
        s, Inches(6.75), Inches(1.15), Inches(6.15), Inches(5.7), SOFT_GREEN,
        "Solution — AdaptiveEdu",
        "A Telegram bot linked to the real 1st Secondary list of subjects.\n\nThe student sends an ID.\nThe bot sends the lesson that is due.\nThen a 3-question quiz.\n\nPass: next lesson.\nFail: same lesson, simpler text.\nStill stuck: the teacher can step in.",
        bar=GREEN, tsize=20, bsize=17,
    )

    # 3 Student journey
    s = blank(prs)
    header(s, "STUDENT JOURNEY", "One path. Easy to demo.", f)
    steps = [
        ("1", "Send ID", "S1-1001\nor\nS1-1001 Integrated Sciences"),
        ("2", "Get lesson", "Only the lesson\nthat is due"),
        ("3", "Take quiz", "Reply with\n3 letters\nC B A"),
        ("4", "See result", "Pass or Fail\nsaved in Sheets"),
        ("5", "Next step", "Pass → next lesson\nFail → retry simple"),
    ]
    for i, (n, t, b) in enumerate(steps):
        x = 0.4 + i * 2.56
        add_round(s, Inches(x), Inches(1.3), Inches(2.4), Inches(5.5), [SOFT_TEAL, SOFT_BLUE, SOFT_GOLD, SOFT_GREEN, SOFT_ORANGE][i])
        add_round(s, Inches(x + 0.85), Inches(1.55), Inches(0.7), Inches(0.7), TEAL)
        add_text(s, Inches(x + 0.85), Inches(1.55), Inches(0.7), Inches(0.7), n, size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, Inches(x + 0.12), Inches(2.45), Inches(2.16), Inches(0.7), t, size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_text(s, Inches(x + 0.12), Inches(3.3), Inches(2.16), Inches(3.1), b, size=15, color=DARK, align=PP_ALIGN.CENTER)

    # 4 Two sheets
    s = blank(prs)
    header(s, "DATA", "Two Google Sheets the bot uses", f)
    pic(s, "adaptiveedu_two_sheets.jpg", Inches(7.55), Inches(1.15), Inches(5.4))
    card(
        s, Inches(0.4), Inches(1.15), Inches(6.9), Inches(2.7), SOFT_BLUE,
        "Sheet 1 — Students",
        "Who is this student?\nID, first name, last name, email, grade.\nn8n writes: Telegram chat, current lesson, last score, Pass/Fail.",
        bar=NAVY2, tsize=18, bsize=15,
    )
    card(
        s, Inches(0.4), Inches(4.05), Inches(6.9), Inches(2.8), SOFT_TEAL,
        "Sheet 2 — Curriculum",
        "Subjects: 1st Secondary list.\nLessons: ordered lessons + quiz answers.\nResults: n8n appends every quiz (never update a row).",
        bar=TEAL, tsize=18, bsize=15,
    )

    # 5 Sheet links + n8n columns
    s = blank(prs)
    header(s, "SHEET LINKS", "Open these in the demo", f)
    card(
        s, Inches(0.4), Inches(1.15), Inches(12.5), Inches(2.15), SOFT_BLUE,
        "Students",
        SHEET_STUDENTS + "\nTab name: Students   ·   Match in n8n on studentID",
        bar=NAVY2, tsize=18, bsize=15,
    )
    card(
        s, Inches(0.4), Inches(3.5), Inches(12.5), Inches(3.35), SOFT_TEAL,
        "Curriculum",
        SHEET_CURR + "\nTabs: Subjects (read)  ·  Lessons (read)  ·  Results (append only)\nPass mark in Lessons = 5. Score is out of 9 (3 questions × 3 points).",
        bar=TEAL, tsize=18, bsize=15,
    )

    # 6 n8n flow
    s = blank(prs)
    header(s, "N8N", "The simplest workflow that still does the whole job", f)
    pic(s, "adaptiveedu_n8n_flow.jpg", Inches(0.35), Inches(1.05), Inches(12.6))
    add_text(
        s, Inches(0.45), Inches(6.35), Inches(12.4), Inches(0.7),
        "Two message types only.  ID  →  send lesson.    Three letters  →  grade quiz.",
        size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER,
    )

    # 7 Nodes
    s = blank(prs)
    header(s, "N8N NODES", "What each node does", f)
    nodes = [
        ("Telegram Trigger", "n8n-nodes-base.telegramTrigger\nNew student message."),
        ("Read message", "n8n-nodes-base.code\nID or quiz answer?"),
        ("IF student ID", "n8n-nodes-base.if\ntrue = ID  ·  false = quiz"),
        ("Get students", "n8n-nodes-base.googleSheets\nRead Students tab."),
        ("Get lessons", "n8n-nodes-base.googleSheets\nRead Lessons tab."),
        ("Pick next lesson", "n8n-nodes-base.code\nDue lesson, or retry if Fail."),
        ("Save open lesson", "n8n-nodes-base.googleSheets\nUpdate current lesson."),
        ("Send lesson + quiz", "n8n-nodes-base.telegram\nLesson text + 3 questions."),
        ("Get students for quiz", "n8n-nodes-base.googleSheets\nFind the student again."),
        ("Get lessons for quiz", "n8n-nodes-base.googleSheets\nLoad answers to grade."),
        ("Grade quiz", "n8n-nodes-base.code\nMatch A–D. Score / 9."),
        ("Save score", "n8n-nodes-base.googleSheets\nUpdate lastScore / lastResult."),
        ("Append result", "n8n-nodes-base.googleSheets\nAppend one Results row."),
        ("Send result", "n8n-nodes-base.telegram\nPass or Fail message."),
    ]
    for i, (t, b) in enumerate(nodes):
        r, c = divmod(i, 7)
        x = Inches(0.28 + c * 1.86)
        y = Inches(1.08 + r * 3.05)
        add_round(s, x, y, Inches(1.76), Inches(2.88), LIGHT if i % 2 == 0 else SOFT_TEAL)
        icon = {
            "Telegram Trigger": "telegram.png",
            "Read message": "code.png",
            "IF student ID": "if.png",
            "Get students": "sheets.png",
            "Get lessons": "sheets.png",
            "Pick next lesson": "code.png",
            "Save open lesson": "sheets.png",
            "Send lesson + quiz": "telegram.png",
            "Get students for quiz": "sheets.png",
            "Get lessons for quiz": "sheets.png",
            "Grade quiz": "code.png",
            "Save score": "sheets.png",
            "Append result": "sheets.png",
            "Send result": "telegram.png",
        }[t]
        ip = DIAG / "n8n_icons" / icon
        if ip.exists():
            s.shapes.add_picture(str(ip), x + Inches(0.58), y + Inches(0.12), Inches(0.58), Inches(0.58))
        add_text(s, x + Inches(0.06), y + Inches(0.78), Inches(1.64), Inches(0.7), t, size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.06), y + Inches(1.48), Inches(1.64), Inches(1.28), b, size=10, color=DARK, align=PP_ALIGN.CENTER)

    # 8 Pass fail
    s = blank(prs)
    header(s, "RULE", "Pass, fail, and the teacher", f)
    card(
        s, Inches(0.4), Inches(1.15), Inches(6.15), Inches(5.7), SOFT_GREEN,
        "Pass  ·  5 or more / 9",
        "The lesson is done.\n\nStudents sheet: lastResult = Pass.\nResults sheet: one new row.\n\nNext time the student sends the ID, the bot opens the next lesson number.",
        bar=GREEN, tsize=20, bsize=17,
    )
    card(
        s, Inches(6.75), Inches(1.15), Inches(6.15), Inches(5.7), SOFT_ORANGE,
        "Fail  ·  below 5",
        "The lesson is not done.\n\nStudents sheet: lastResult = Fail.\nThe lesson number does not go up.\n\nNext ID message: same lesson, simpler text.\nIf still stuck, a teacher can help by hand.",
        bar=ORANGE, tsize=20, bsize=17,
    )

    # 9 Tools
    s = blank(prs)
    header(s, "TOOLS", "Small stack. Easy to explain.", f)
    tools = [
        ("Telegram", "The only student screen."),
        ("n8n", "Reads sheets, grades, sends."),
        ("Students sheet", "Who the student is now."),
        ("Curriculum sheet", "What to teach + quiz log."),
        ("Code nodes", "Pick lesson. Grade A–D."),
        ("GitHub", GITHUB.replace("https://", "")),
    ]
    for i, (t, b) in enumerate(tools):
        r, c = divmod(i, 3)
        card(s, Inches(0.4 + c * 4.22), Inches(1.15 + r * 2.9), Inches(4.05), Inches(2.7), [SOFT_TEAL, SOFT_BLUE, SOFT_GOLD][c], t, b, bar=TEAL, tsize=18, bsize=16)

    # 10 Demo script
    s = blank(prs)
    header(s, "LIVE DEMO", "Say this. Do this. 2 minutes.", f)
    demo = [
        ("1", "Open Telegram", "Send: S1-1001"),
        ("2", "Show the bot reply", "Lesson 1 text + 3 questions"),
        ("3", "Send answers", "C B A  (or any 3 letters)"),
        ("4", "Open Students sheet", "lastScore and lastResult filled"),
        ("5", "Open Results tab", "A new row was appended"),
        ("6", "If Fail", "Send S1-1001 again → same lesson, simpler"),
    ]
    for i, (n, t, b) in enumerate(demo):
        y = 1.15 + i * 0.95
        add_round(s, Inches(0.4), Inches(y), Inches(12.5), Inches(0.88), SOFT_TEAL if i % 2 == 0 else LIGHT)
        add_round(s, Inches(0.55), Inches(y + 0.16), Inches(0.55), Inches(0.55), TEAL)
        add_text(s, Inches(0.55), Inches(y + 0.16), Inches(0.55), Inches(0.55), n, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, Inches(1.35), Inches(y + 0.12), Inches(4.3), Inches(0.64), t, size=18, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, Inches(5.8), Inches(y + 0.12), Inches(6.8), Inches(0.64), b, size=16, color=DARK, anchor=MSO_ANCHOR.MIDDLE)

    # 11 Weeks
    s = blank(prs)
    header(s, "PLAN", "Two weeks to delivery", f)
    pic(s, "adaptiveedu_milestones.jpg", Inches(0.4), Inches(1.05), Inches(12.5))
    miles = [
        ("Week 1", "Sheets live.\nImport n8n JSON.\nConnect Telegram + Sheets.\nDemo ID → lesson."),
        ("Week 2", "Grade quiz + Pass/Fail.\nTest 3 IDs.\nVideo + this PPT.\nPush GitHub and submit."),
    ]
    for i, (t, b) in enumerate(miles):
        add_text(s, Inches(1.4 + i * 6.3), Inches(5.4), Inches(4.4), Inches(1.55), f"{t}\n{b}", size=16, color=DARK, align=PP_ALIGN.CENTER)

    # 12 Deliverables
    s = blank(prs)
    header(s, "SUBMIT", "What the panel should see", f)
    dels = [
        ("Working bot", "Ask ID → lesson → quiz → sheet update."),
        ("n8n JSON", "AdaptiveEdu_Telegram_Bot.json"),
        ("Two Sheets", "Students + Curriculum links."),
        ("This PPT", "Any teammate can present it."),
        ("GitHub", "https://github.com/ahmedhamedali82/AdaptiveEdu"),
        ("Video 2–5 min", "Problem → demo → Pass/Fail."),
    ]
    for i, (t, b) in enumerate(dels):
        r, c = divmod(i, 3)
        card(s, Inches(0.4 + c * 4.22), Inches(1.15 + r * 2.9), Inches(4.05), Inches(2.7), LIGHT, t, b, bar=TEAL, tsize=18, bsize=16)

    prs.save(OUT)
    print(OUT)
    print("slides", len(prs.slides))


if __name__ == "__main__":
    build()
